import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "data"))

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

from trade_data import load_trades

FONT = "Arial"
OUT_DIR = os.environ.get("DASHBOARD_OUT_DIR", os.path.join(os.path.dirname(__file__), "..", "..", "_build"))
PATH = os.path.join(OUT_DIR, "dashboard_step3.xlsx")
wb = load_workbook(PATH)

FIRST_DATA_ROW = 5
N_TRADES = len(load_trades())
LAST_ROW = FIRST_DATA_ROW + N_TRADES - 1 + 60

header_font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="1F4E78")
title_font = Font(name=FONT, bold=True, size=16, color="1F4E78")
sub_font = Font(name=FONT, italic=True, size=10, color="666666")
normal_font = Font(name=FONT, size=11)
bold_font = Font(name=FONT, size=11, bold=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
CUR0 = '$#,##0;[RED]($#,##0)'
PCT = '0.0%'

ws = wb.create_sheet("Confidence Calibration")
ws.sheet_view.showGridLines = False

ws["A1"] = "Confidence Calibration"
ws["A1"].font = title_font
ws["A2"] = ("Checks whether the confidence % given on each call actually tracks the real win rate. "
            "Trades that predate the confidence-scoring format, or were off-script entries, are excluded automatically (blank confidence).")
ws["A2"].font = sub_font
ws.merge_cells("A2:G2")

widths = [22, 10, 8, 10, 14, 16, 40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

headers = ["Confidence Band", "Trades", "Wins", "Win Rate", "Total P&L", "Avg P&L / Trade", "Read"]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=4, column=i, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = center
    c.border = border

M = f"'Trade Log'!$M${FIRST_DATA_ROW}:$M${LAST_ROW}"
K = f"'Trade Log'!$K${FIRST_DATA_ROW}:$K${LAST_ROW}"

bands = [
    ("Under 40%", 0, 40),
    ("40% - 49%", 40, 50),
    ("50% - 59%", 50, 60),
    ("60% - 69%", 60, 70),
    ("70% - 79%", 70, 80),
    ("80%+", 80, 101),
]

start_row = 5
for i, (label, lo, hi) in enumerate(bands):
    r = start_row + i
    ws.cell(row=r, column=1, value=label).font = normal_font
    ws.cell(row=r, column=2,
            value=f'=COUNTIFS({M},">="&{lo},{M},"<"&{hi})')
    ws.cell(row=r, column=3,
            value=f'=COUNTIFS({M},">="&{lo},{M},"<"&{hi},{K},">0")')
    ws.cell(row=r, column=4, value=f'=IFERROR(C{r}/B{r},"")')
    ws.cell(row=r, column=4).number_format = PCT
    ws.cell(row=r, column=5,
            value=f'=SUMIFS({K},{M},">="&{lo},{M},"<"&{hi})')
    ws.cell(row=r, column=5).number_format = CUR0
    ws.cell(row=r, column=6, value=f'=IFERROR(E{r}/B{r},"")')
    ws.cell(row=r, column=6).number_format = CUR0
    ws.cell(row=r, column=7,
            value=(f'=IF(B{r}=0,"no trades yet",'
                    f'IF(D{r}>=({lo}+{hi})/2/100,"tracking or beating this band","underperforming this band"))'))
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = border
        if c != 7:
            ws.cell(row=r, column=c).font = normal_font
        else:
            ws.cell(row=r, column=c).font = normal_font
            ws.cell(row=r, column=c).alignment = left

total_row = start_row + len(bands)
ws.cell(row=total_row, column=1, value="ALL SCORED TRADES").font = bold_font
ws.cell(row=total_row, column=2, value=f"=SUM(B{start_row}:B{total_row-1})").font = bold_font
ws.cell(row=total_row, column=3, value=f"=SUM(C{start_row}:C{total_row-1})").font = bold_font
ws.cell(row=total_row, column=4, value=f"=IFERROR(C{total_row}/B{total_row},\"\")")
ws.cell(row=total_row, column=4).number_format = PCT
ws.cell(row=total_row, column=4).font = bold_font
ws.cell(row=total_row, column=5, value=f"=SUM(E{start_row}:E{total_row-1})")
ws.cell(row=total_row, column=5).number_format = CUR0
ws.cell(row=total_row, column=5).font = bold_font
for c in range(1, 8):
    ws.cell(row=total_row, column=c).border = Border(top=Side(style="double"), left=thin, right=thin, bottom=thin)

green_fill = PatternFill("solid", fgColor="C6EFCE")
red_fill = PatternFill("solid", fgColor="FFC7CE")
rng = f"E{start_row}:E{total_row}"
ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=green_fill))
ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=red_fill))

chart = BarChart()
chart.type = "col"
chart.title = "Win Rate by Confidence Band"
chart.y_axis.title = "Win Rate"
chart.style = 10
data = Reference(ws, min_col=4, min_row=4, max_row=start_row + len(bands) - 1)
cats = Reference(ws, min_col=1, min_row=start_row, max_row=start_row + len(bands) - 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.legend = None
chart.height = 8
chart.width = 16
ws.add_chart(chart, f"A{total_row + 3}")

note_row = total_row + 22
ws.cell(row=note_row, column=1,
        value=("How to use this before the next call: check this table (and the By Strategy / By Session tabs) "
               "first. A band or category with a win rate meaningfully below the confidence number being considered "
               "means that confidence should be scored lower this time; a band consistently beating its number "
               "justifies scoring it higher. This is a manual calibration process. The separate ml_train.py script "
               "(funded-trading-assistant/data/) runs a real logistic-regression model on the same trades.csv for "
               "a second, model-based read alongside this table."))
ws.cell(row=note_row, column=1).font = sub_font
ws.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells(f"A{note_row}:G{note_row+3}")

wb.save(os.path.join(OUT_DIR, "dashboard_step4.xlsx"))
print("Step 4 saved. Sheets:", wb.sheetnames)
