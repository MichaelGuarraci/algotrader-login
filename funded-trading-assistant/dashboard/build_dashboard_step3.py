import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "data"))

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.utils import get_column_letter

from trade_data import load_trades

FONT = "Arial"
OUT_DIR = os.environ.get("DASHBOARD_OUT_DIR", os.path.join(os.path.dirname(__file__), "..", "..", "_build"))
PATH = os.path.join(OUT_DIR, "dashboard_step2.xlsx")
wb = load_workbook(PATH)
tl = wb["Trade Log"]
dash = wb["Dashboard"]

FIRST_DATA_ROW = 5
N_TRADES = len(load_trades())
LAST_ROW = FIRST_DATA_ROW + N_TRADES - 1 + 60

header_font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor="1F4E78")
title_font = Font(name=FONT, bold=True, size=16, color="1F4E78")
sub_font = Font(name=FONT, italic=True, size=10, color="666666")
section_font = Font(name=FONT, bold=True, size=12, color="1F4E78")
normal_font = Font(name=FONT, size=11)
bold_font = Font(name=FONT, size=11, bold=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center")
CUR0 = '$#,##0;[RED]($#,##0)'
PCT = '0.0%'

def make_breakdown_sheet(name, categories, group_col_letter, title):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A2"] = "Breaks down every closed trade by " + ("strategy type." if group_col_letter == "L" else "session/time-of-day.")
    ws["A2"].font = sub_font

    headers = ["Category", "Trades", "Wins", "Win Rate", "Total P&L", "Avg P&L / Trade", "% of Total P&L"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border
    widths = [28, 10, 8, 10, 14, 16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    grp = f"'Trade Log'!${group_col_letter}${FIRST_DATA_ROW}:${group_col_letter}${LAST_ROW}"
    pnl = f"'Trade Log'!$K${FIRST_DATA_ROW}:$K${LAST_ROW}"

    start_row = 5
    for i, cat in enumerate(categories):
        r = start_row + i
        ws.cell(row=r, column=1, value=cat).font = normal_font
        ws.cell(row=r, column=2, value=f'=COUNTIF({grp},$A{r})')
        ws.cell(row=r, column=3, value=f'=COUNTIFS({grp},$A{r},{pnl},">0")')
        ws.cell(row=r, column=4, value=f'=IFERROR(C{r}/B{r},"")')
        ws.cell(row=r, column=4).number_format = PCT
        ws.cell(row=r, column=5, value=f'=SUMIF({grp},$A{r},{pnl})')
        ws.cell(row=r, column=5).number_format = CUR0
        ws.cell(row=r, column=6, value=f'=IFERROR(E{r}/B{r},"")')
        ws.cell(row=r, column=6).number_format = CUR0
        for c in range(1, 8):
            ws.cell(row=r, column=c).border = border
            if c not in (4,):
                ws.cell(row=r, column=c).font = normal_font

    total_row = start_row + len(categories)
    ws.cell(row=total_row, column=1, value="TOTAL").font = bold_font
    ws.cell(row=total_row, column=2, value=f"=SUM(B{start_row}:B{total_row-1})").font = bold_font
    ws.cell(row=total_row, column=3, value=f"=SUM(C{start_row}:C{total_row-1})").font = bold_font
    ws.cell(row=total_row, column=4, value=f"=IFERROR(C{total_row}/B{total_row},\"\")")
    ws.cell(row=total_row, column=4).number_format = PCT
    ws.cell(row=total_row, column=4).font = bold_font
    ws.cell(row=total_row, column=5, value=f"=SUM(E{start_row}:E{total_row-1})")
    ws.cell(row=total_row, column=5).number_format = CUR0
    ws.cell(row=total_row, column=5).font = bold_font
    ws.cell(row=total_row, column=7, value=1).number_format = PCT
    ws.cell(row=total_row, column=7).font = bold_font
    for c in range(1, 8):
        ws.cell(row=total_row, column=c).border = Border(top=Side(style="double"), left=thin, right=thin, bottom=thin)

    # % of total P&L column for category rows (references total row's total P&L)
    for i in range(len(categories)):
        r = start_row + i
        ws.cell(row=r, column=7, value=f'=IFERROR(E{r}/$E${total_row},"")')
        ws.cell(row=r, column=7).number_format = PCT
        ws.cell(row=r, column=7).font = normal_font
        ws.cell(row=r, column=7).border = border

    # conditional format total P&L col
    from openpyxl.formatting.rule import CellIsRule
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    rng = f"E{start_row}:E{total_row}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=green_fill))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=red_fill))

    # bar chart: Total P&L by category
    chart = BarChart()
    chart.type = "col"
    chart.title = title.replace(" Breakdown", "") + " — Total P&L"
    chart.y_axis.title = "P&L ($)"
    chart.x_axis.title = None
    chart.style = 10
    data = Reference(ws, min_col=5, min_row=4, max_row=start_row + len(categories) - 1)
    cats = Reference(ws, min_col=1, min_row=start_row, max_row=start_row + len(categories) - 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.height = 8
    chart.width = 18
    ws.add_chart(chart, f"A{total_row + 3}")
    return ws, start_row, total_row

strategy_categories = ["Trend Continuation", "Pullback / Retest Entry", "Resistance / Support Fade",
                        "Breakout", "Counter-Trend Bounce", "Unmanaged Overnight Hold"]
session_categories = ["Pre-Market", "RTH", "Overnight/Globex", "Weekend Reopen"]

make_breakdown_sheet("By Strategy", strategy_categories, "L", "Performance by Strategy Type")
make_breakdown_sheet("By Session", session_categories, "D", "Performance by Session")

# =========================================================
# Equity curve chart on Dashboard
# =========================================================
chart = LineChart()
chart.title = "Equity Curve"
chart.style = 12
chart.y_axis.title = "Balance ($)"
chart.x_axis.title = "Trade #"
data = Reference(tl, min_col=16, min_row=4, max_row=LAST_ROW)  # column P, include header row 4
cats = Reference(tl, min_col=1, min_row=FIRST_DATA_ROW, max_row=LAST_ROW)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
s = chart.series[0]
s.marker = Marker(symbol="circle", size=5)
s.smooth = False
chart.height = 9
chart.width = 20
chart.legend = None
dash.add_chart(chart, "A27")

wb.save(os.path.join(OUT_DIR, "dashboard_step3.xlsx"))
print("Step 3 saved. Sheets:", wb.sheetnames)
